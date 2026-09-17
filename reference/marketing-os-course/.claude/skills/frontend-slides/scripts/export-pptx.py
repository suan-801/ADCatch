#!/usr/bin/env python3
"""HTML 프레젠테이션을 PowerPoint(.pptx) 파일로 export.

기본 모드 (--editable, 기본값):
    하이브리드 방식. 텍스트만 숨긴 슬라이드 스크린샷을 배경으로 깔고,
    각 텍스트 노드를 native PPT 텍스트박스로 위에 오버레이.
    → PowerPoint에서 텍스트 클릭·편집 가능, 그라디언트·도형 비주얼 보존.

이미지 모드 (--image-only):
    각 슬라이드 전체를 단일 이미지로 임베드. 텍스트 편집 불가.
    복잡한 비주얼이 압도적이고 편집이 필요 없을 때.

Usage:
    python3 export-pptx.py <input.html> [output.pptx] [--editable | --image-only] [--compact]

Examples:
    python3 scripts/export-pptx.py ./presentation.html                     # 편집 가능 모드 (기본)
    python3 scripts/export-pptx.py ./presentation.html ./deck.pptx
    python3 scripts/export-pptx.py ./presentation.html --image-only       # 이미지 모드
    python3 scripts/export-pptx.py ./presentation.html --compact          # 1280x720

Requires:
    pip install playwright python-pptx
    python -m playwright install chromium
"""

from __future__ import annotations

import argparse
import http.server
import re
import socket
import socketserver
import sys
import tempfile
import threading
import time
from contextlib import contextmanager
from pathlib import Path


def _eprint(*args, **kwargs):
    kwargs.setdefault("file", sys.stderr)
    print(*args, **kwargs)


def _find_free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@contextmanager
def serve_directory(directory: Path):
    """`directory`를 HTTP로 서빙하는 백그라운드 서버. 컨텍스트 종료 시 자동 종료."""
    port = _find_free_port()

    class _Handler(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, directory=str(directory), **kwargs)

        def log_message(self, format, *args):  # 로그 끄기
            pass

    httpd = socketserver.TCPServer(("127.0.0.1", port), _Handler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    try:
        yield port
    finally:
        httpd.shutdown()
        httpd.server_close()


_TEXT_LAYER_JS = r"""
(slideIdx) => {
    const slides = document.querySelectorAll('.slide');
    const slide = slides[slideIdx];
    const slideRect = slide.getBoundingClientRect();
    const result = [];

    function visible(el, cs) {
        if (cs.display === 'none' || cs.visibility === 'hidden') return false;
        if (parseFloat(cs.opacity) < 0.05) return false;
        const r = el.getBoundingClientRect();
        if (r.width < 1 || r.height < 1) return false;
        return true;
    }

    function hasOnlyTextOrBr(el) {
        // True if children are only text nodes or <br> elements (no other text-bearing elements)
        for (const c of el.childNodes) {
            if (c.nodeType === 3) continue;             // text node
            if (c.nodeType === 1) {
                if (c.tagName === 'BR') continue;
                if (c.textContent && c.textContent.trim()) return false;
            }
        }
        return true;
    }

    const walker = document.createTreeWalker(slide, NodeFilter.SHOW_ELEMENT);
    let node;
    while (node = walker.nextNode()) {
        const text = (node.innerText || '').trim();
        if (!text) continue;
        if (!hasOnlyTextOrBr(node)) continue;
        const cs = getComputedStyle(node);
        if (!visible(node, cs)) continue;

        const rect = node.getBoundingClientRect();
        result.push({
            text: node.innerText,
            x: rect.left - slideRect.left,
            y: rect.top - slideRect.top,
            w: rect.width,
            h: rect.height,
            fontFamily: cs.fontFamily,
            fontSize: parseFloat(cs.fontSize),
            fontWeight: cs.fontWeight,
            fontStyle: cs.fontStyle,
            color: cs.color,
            textAlign: cs.textAlign,
            lineHeight: cs.lineHeight,
            letterSpacing: cs.letterSpacing,
        });
    }
    return result;
}
"""


_HIDE_TEXT_CSS = """
.slide *, .slide *::before, .slide *::after {
    color: transparent !important;
    -webkit-text-fill-color: transparent !important;
    text-shadow: none !important;
    caret-color: transparent !important;
}
"""


def capture_slides(
    html_path: Path,
    viewport: tuple[int, int],
    out_dir: Path,
    editable: bool,
) -> tuple[list[Path], list[list[dict]]]:
    """HTML을 띄우고 각 .slide 섹션 스크린샷 + (옵션) 텍스트 레이어 추출.

    Returns:
        (screenshot_paths, text_layers): 슬라이드별 PNG 경로 / 텍스트 노드 리스트.
        editable=False면 text_layers는 빈 리스트들.
    """
    from playwright.sync_api import sync_playwright

    html_path = html_path.resolve()
    serve_dir = html_path.parent
    rel_url = html_path.name

    out_dir.mkdir(parents=True, exist_ok=True)
    screenshot_paths: list[Path] = []
    text_layers: list[list[dict]] = []

    with serve_directory(serve_dir) as port, sync_playwright() as p:
        url = f"http://127.0.0.1:{port}/{rel_url}"
        browser = p.chromium.launch()
        context = browser.new_context(
            viewport={"width": viewport[0], "height": viewport[1]},
            device_scale_factor=2,
        )
        page = context.new_page()
        page.goto(url, wait_until="networkidle", timeout=30000)

        try:
            page.evaluate("document.fonts.ready")
        except Exception:
            time.sleep(0.5)

        page.add_style_tag(content="""
            *, *::before, *::after {
                animation: none !important;
                transition: none !important;
                scroll-behavior: auto !important;
            }
        """)

        slide_count = page.evaluate("document.querySelectorAll('.slide').length")
        if slide_count == 0:
            raise RuntimeError("HTML에서 .slide 섹션을 찾지 못했습니다.")
        _eprint(f"  → {slide_count}개 슬라이드 발견")

        # Editable 모드: 먼저 텍스트 레이어를 (스타일 변경 전에) 모두 추출
        if editable:
            _eprint("  → 텍스트 레이어 추출 중…")
            for i in range(slide_count):
                page.evaluate(f"""
                    () => {{
                        const slides = document.querySelectorAll('.slide');
                        slides[{i}].scrollIntoView({{behavior: 'instant', block: 'start'}});
                    }}
                """)
                page.wait_for_timeout(150)
                nodes = page.evaluate(_TEXT_LAYER_JS, i)
                text_layers.append(nodes)
                _eprint(f"    [{i + 1:02d}/{slide_count}] {len(nodes)}개 텍스트 노드")

            # 텍스트 숨기고 → 배경만 캡처
            page.add_style_tag(content=_HIDE_TEXT_CSS)
            page.wait_for_timeout(200)
        else:
            text_layers = [[] for _ in range(slide_count)]

        for i in range(slide_count):
            page.evaluate(f"""
                () => {{
                    const slides = document.querySelectorAll('.slide');
                    slides[{i}].scrollIntoView({{behavior: 'instant', block: 'start'}});
                }}
            """)
            page.wait_for_timeout(300)

            shot = out_dir / f"slide_{i + 1:02d}.png"
            slide_handle = page.query_selector_all(".slide")[i]
            slide_handle.screenshot(path=str(shot), type="png")
            screenshot_paths.append(shot)
            _eprint(f"    [{i + 1:02d}/{slide_count}] {shot.name}")

        browser.close()

    return screenshot_paths, text_layers


_RGB_RE = re.compile(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)")


def _parse_color(css: str):
    m = _RGB_RE.match(css or "")
    if not m:
        return None
    return int(m.group(1)), int(m.group(2)), int(m.group(3))


def _is_bold(weight: str) -> bool:
    if not weight:
        return False
    if weight.isdigit():
        return int(weight) >= 600
    return weight.lower() in ("bold", "bolder")


_ALIGN_MAP_KEY = {"left": "left", "center": "center", "right": "right", "justify": "justify",
                  "start": "left", "end": "right"}


def build_pptx(
    screenshots: list[Path],
    text_layers: list[list[dict]],
    output: Path,
    viewport: tuple[int, int],
):
    """배경 PNG + (옵션) 편집 가능 텍스트 오버레이로 .pptx 생성."""
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    prs = Presentation()
    # 16:9 (13.333" × 7.5") = 12192000 × 6858000 EMU
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # CSS px → EMU 환산: 뷰포트 폭이 슬라이드 폭에 매핑됨
    px_to_emu = prs.slide_width / viewport[0]

    blank_layout = prs.slide_layouts[6]
    align_enum = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }

    for shot, nodes in zip(screenshots, text_layers):
        slide = prs.slides.add_slide(blank_layout)
        # 배경 풀블리드 이미지
        slide.shapes.add_picture(
            str(shot),
            left=0,
            top=0,
            width=prs.slide_width,
            height=prs.slide_height,
        )

        # 편집 가능 텍스트 오버레이
        for n in nodes:
            text = (n.get("text") or "").strip("\n")
            if not text:
                continue

            # 살짝 여유 — 짤림 방지용 margin 보정
            pad_x_emu = int(2 * px_to_emu)
            pad_y_emu = int(2 * px_to_emu)
            left = max(0, int(n["x"] * px_to_emu) - pad_x_emu)
            top = max(0, int(n["y"] * px_to_emu) - pad_y_emu)
            width = max(int(0.5 * px_to_emu * 10), int(n["w"] * px_to_emu) + 2 * pad_x_emu)
            height = max(int(0.5 * px_to_emu * 10), int(n["h"] * px_to_emu) + 2 * pad_y_emu)

            tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = Emu(0)
            tf.margin_right = Emu(0)
            tf.margin_top = Emu(0)
            tf.margin_bottom = Emu(0)
            tf.vertical_anchor = MSO_ANCHOR.TOP

            family_raw = (n.get("fontFamily") or "").split(",")[0].strip().strip('"\'')
            color_rgb = _parse_color(n.get("color", ""))
            font_size_pt = max(6.0, n["fontSize"] * 0.75)  # CSS px → pt (1px ≈ 0.75pt)
            bold = _is_bold(n.get("fontWeight", ""))
            italic = n.get("fontStyle") == "italic"
            align_key = _ALIGN_MAP_KEY.get(n.get("textAlign", "left"), "left")
            align = align_enum.get(align_key, PP_ALIGN.LEFT)

            # 줄바꿈 → 단락 분리 (innerText는 \n으로 줄바꿈)
            lines = text.split("\n")
            for li, line in enumerate(lines):
                p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                p.alignment = align
                run = p.add_run()
                run.text = line
                f = run.font
                if family_raw:
                    f.name = family_raw
                f.size = Pt(font_size_pt)
                f.bold = bold
                f.italic = italic
                if color_rgb:
                    f.color.rgb = RGBColor(*color_rgb)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def main():
    ap = argparse.ArgumentParser(description="Export HTML presentation to PowerPoint (.pptx)")
    ap.add_argument("input", help="Path to input HTML file")
    ap.add_argument("output", nargs="?", help="Path to output .pptx (default: alongside input)")
    ap.add_argument("--compact", action="store_true", help="1280x720 viewport (smaller file)")
    mode = ap.add_mutually_exclusive_group()
    mode.add_argument("--editable", dest="editable", action="store_true",
                      help="텍스트 편집 가능 모드 (기본). 그라디언트 배경+native 텍스트박스 하이브리드.")
    mode.add_argument("--image-only", dest="editable", action="store_false",
                      help="이미지 전용 모드. 각 슬라이드 = 단일 이미지, 텍스트 편집 불가.")
    ap.set_defaults(editable=True)
    args = ap.parse_args()

    html_path = Path(args.input).resolve()
    if not html_path.exists():
        _eprint(f"✗ HTML 파일을 찾을 수 없습니다: {html_path}")
        sys.exit(1)

    if args.output:
        output = Path(args.output).resolve()
    else:
        output = html_path.with_suffix(".pptx")

    viewport = (1280, 720) if args.compact else (1920, 1080)
    mode_label = "editable (하이브리드)" if args.editable else "image-only"

    _eprint(f"ℹ 입력  : {html_path}")
    _eprint(f"ℹ 출력  : {output}")
    _eprint(f"ℹ 해상도: {viewport[0]}×{viewport[1]} ({'compact' if args.compact else 'full HD'})")
    _eprint(f"ℹ 모드  : {mode_label}")

    with tempfile.TemporaryDirectory(prefix="html2pptx_") as tmp:
        tmp_dir = Path(tmp)
        _eprint("ℹ 슬라이드 캡처 중…")
        screenshots, text_layers = capture_slides(html_path, viewport, tmp_dir, args.editable)

        _eprint("ℹ .pptx 빌드 중…")
        build_pptx(screenshots, text_layers, output, viewport)

    size_kb = output.stat().st_size / 1024
    total_text = sum(len(layer) for layer in text_layers)
    extra = f", 텍스트 노드 {total_text}개 편집 가능" if args.editable else ""
    _eprint(f"✓ 완료: {output} ({size_kb:.1f} KB, {len(screenshots)} 슬라이드{extra})")


if __name__ == "__main__":
    main()
